"""Recompute model outputs from current inputs.

This module is the deterministic counterpart to a spreadsheet recalc engine.
For each supported model schema, it knows how to read the input cells, run
the arithmetic, and write the output cells.

Why we don't rely on Excel formulas alone
-----------------------------------------
openpyxl writes formula text but does not evaluate it. Some teams have
LibreOffice or Excel available for headless recalc; many do not. By doing
the recalc in Python we keep the CLI independent of any office suite, and
we get a single auditable place where the model's logic lives.

The schema is named in `deal.yaml > compute_method`. Two are supported:

  - `availability_payment_pf`   — Highway 407 East Extension style
  - `merchant_capacity_pf`      — Oneida Energy Storage style

To onboard a new deal that doesn't fit either, add a `compute_method` here.
"""

from __future__ import annotations

from dataclasses import dataclass
from typing import Callable

import openpyxl
from openpyxl.utils import get_column_letter

from .deal import Deal
from .paths import deal_model_path


@dataclass
class ComputeReport:
    method: str
    output_writes: dict[str, float]
    notes: list[str]


def _read_quarterly(ws, base_row: int, count: int = 8) -> list[float]:
    return [float(ws.cell(row=base_row + i, column=2).value or 0) for i in range(count)]


def _availability_payment_pf(deal: Deal, wb) -> ComputeReport:
    """Highway-407 style: revenue = availability payment, no arbitrage component."""
    inp = wb["Inputs"]
    ds = wb["DebtSchedule"]
    cf = wb["CashFlow"]
    out = wb["Outputs"]

    # Inputs layout (column B):
    # rows 2..9   AvailabilityPayment_Q1..Q8
    # rows 10..17 Opex_Q1..Q8
    # row 18      InterestRate
    # row 19      InflationIndex
    # row 20      DSRA_Target
    # row 21      DSRA_Balance
    revenue = _read_quarterly(inp, 2)
    opex = _read_quarterly(inp, 10)
    rate = float(inp["B18"].value)
    dsra_balance = float(inp["B21"].value)

    return _project_finance_compute(
        ds=ds, cf=cf, out=out,
        revenue=revenue, arbitrage=[0.0] * 8, opex=opex, rate=rate,
        dsra_balance=dsra_balance,
        opening_balance=432_000_000.0,
        amort=4_500_000.0,
        method="availability_payment_pf",
    )


def _merchant_capacity_pf(deal: Deal, wb) -> ComputeReport:
    """Oneida-style: capacity payment + arbitrage."""
    inp = wb["Inputs"]
    ds = wb["DebtSchedule"]
    cf = wb["CashFlow"]
    out = wb["Outputs"]

    # Inputs layout (column B):
    # rows 2..9   CapacityPayment_Q1..Q8
    # rows 10..17 ArbitrageRevenue_Q1..Q8
    # rows 18..25 Opex_Q1..Q8
    # row 26      InterestRate
    # row 27      DSRA_Target
    # row 28      DSRA_Balance
    capacity = _read_quarterly(inp, 2)
    arbitrage = _read_quarterly(inp, 10)
    opex = _read_quarterly(inp, 18)
    rate = float(inp["B26"].value)
    dsra_balance = float(inp["B28"].value)

    return _project_finance_compute(
        ds=ds, cf=cf, out=out,
        revenue=capacity, arbitrage=arbitrage, opex=opex, rate=rate,
        dsra_balance=dsra_balance,
        opening_balance=162_000_000.0,
        amort=2_350_000.0,
        method="merchant_capacity_pf",
    )


def _project_finance_compute(
    *,
    ds, cf, out,
    revenue: list[float],
    arbitrage: list[float],
    opex: list[float],
    rate: float,
    dsra_balance: float,
    opening_balance: float,
    amort: float,
    method: str,
) -> ComputeReport:
    notes: list[str] = []

    # DebtSchedule
    open_bal = [opening_balance]
    interest = []
    close_bal = []
    debt_service = []
    for q in range(8):
        i = open_bal[q] * rate / 4
        c = open_bal[q] - amort
        interest.append(i)
        close_bal.append(c)
        debt_service.append(amort + i)
        if q + 1 < 8:
            open_bal.append(c)

    # Write DebtSchedule rows (row 2 opening, 3 amort, 4 interest, 5 closing, 6 debt service)
    for q in range(8):
        col = q + 2
        ds.cell(row=2, column=col).value = open_bal[q]
        ds.cell(row=3, column=col).value = amort
        ds.cell(row=4, column=col).value = interest[q]
        ds.cell(row=5, column=col).value = close_bal[q]
        ds.cell(row=6, column=col).value = debt_service[q]

    # CashFlow — sheet shape differs slightly between the two schemas. Detect:
    # row 2 = revenue (and capacity for oneida); row 3 = opex (407) or arbitrage (oneida).
    cfads = [revenue[q] + arbitrage[q] - opex[q] for q in range(8)]
    if arbitrage and any(a != 0 for a in arbitrage):
        # Oneida shape: row 2 capacity, row 3 arbitrage, row 4 opex, row 5 cfads,
        # row 6 debtservice, row 7 dscr.
        for q in range(8):
            col = q + 2
            cf.cell(row=2, column=col).value = revenue[q]
            cf.cell(row=3, column=col).value = arbitrage[q]
            cf.cell(row=4, column=col).value = opex[q]
            cf.cell(row=5, column=col).value = cfads[q]
            cf.cell(row=6, column=col).value = debt_service[q]
            cf.cell(row=7, column=col).value = (cfads[q] / debt_service[q]) if debt_service[q] else 0
        dscr_series = [cf.cell(row=7, column=q + 2).value for q in range(8)]
        cfads_row = 5
    else:
        # 407 shape: row 2 revenue, row 3 opex, row 4 cfads, row 5 debtservice, row 6 dscr.
        for q in range(8):
            col = q + 2
            cf.cell(row=2, column=col).value = revenue[q]
            cf.cell(row=3, column=col).value = opex[q]
            cf.cell(row=4, column=col).value = cfads[q]
            cf.cell(row=5, column=col).value = debt_service[q]
            cf.cell(row=6, column=col).value = (cfads[q] / debt_service[q]) if debt_service[q] else 0
        dscr_series = [cf.cell(row=6, column=q + 2).value for q in range(8)]
        cfads_row = 4

    period_dscr = dscr_series[0]
    min_dscr = min(dscr_series) if dscr_series else 0.0

    # LLCR ≈ NPV of remaining CFADS over loan life / opening senior debt.
    # The model only carries 8 quarters; we extrapolate the average annualized
    # CFADS as a level annuity over the remaining tenor (closed-form).
    # `remaining_quarters` is approximate — for the demo deals it's set per
    # method below.
    remaining_quarters = _remaining_quarters_by_method.get(method, 80)
    annualized_cfads = sum(cfads) / len(cfads) * 4
    annual_rate = rate
    if annual_rate > 0:
        annuity_factor = (1 - (1 + annual_rate) ** -(remaining_quarters / 4)) / annual_rate
    else:
        annuity_factor = remaining_quarters / 4
    llcr = (annualized_cfads * annuity_factor) / opening_balance

    out["B3"].value = period_dscr
    out["B4"].value = min_dscr
    out["B5"].value = llcr
    out["B6"].value = close_bal[0]
    out["B7"].value = dsra_balance
    out["B8"].value = _dsra_target_by_method.get(method, dsra_balance)
    out["B9"].value = 0.0  # demo deals are fully drawn at COD
    out["B10"].value = debt_service[0]
    out["B11"].value = cfads[0]

    return ComputeReport(
        method=method,
        output_writes={
            "DSCR": period_dscr,
            "MinDSCRForecast": min_dscr,
            "LLCR": llcr,
            "DebtBalance": close_bal[0],
            "DSRA": dsra_balance,
            "DebtService": debt_service[0],
            "CFADS": cfads[0],
        },
        notes=notes,
    )


_METHODS: dict[str, Callable[[Deal, "openpyxl.Workbook"], ComputeReport]] = {
    "availability_payment_pf": _availability_payment_pf,
    "merchant_capacity_pf":    _merchant_capacity_pf,
}

# Remaining loan life (in quarters) used for LLCR extrapolation. These are
# specific to the demo deals; a real model would compute this from the
# repayment schedule on each run.
_remaining_quarters_by_method = {
    "availability_payment_pf": 88,   # ~22 yrs left on the 25-yr Highway 407 facility
    "merchant_capacity_pf":    64,   # ~16 yrs left on the Oneida facility
}

_dsra_target_by_method = {
    "availability_payment_pf": 18_000_000.0,
    "merchant_capacity_pf":     6_500_000.0,
}


def recompute(deal: Deal, method: str | None = None) -> ComputeReport:
    """Open the deal's model, recompute, save."""
    method = method or _infer_method(deal)
    fn = _METHODS.get(method)
    if fn is None:
        raise ValueError(f"Unknown compute_method {method!r}. Known: {list(_METHODS)}.")
    path = deal_model_path(deal.id)
    wb = openpyxl.load_workbook(path)
    try:
        report = fn(deal, wb)
        wb.save(path)
    finally:
        wb.close()
    return report


def _infer_method(deal: Deal) -> str:
    """Best-effort fallback so deal.yaml doesn't have to declare it."""
    if any(k.startswith("AvailabilityPayment_") for k in deal.input_map):
        return "availability_payment_pf"
    if any(k.startswith("CapacityPayment_") for k in deal.input_map):
        return "merchant_capacity_pf"
    raise ValueError(
        f"Cannot infer compute_method for deal {deal.id!r}. "
        "Declare compute_method in deal.yaml."
    )
