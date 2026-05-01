"""Excel read/write helpers. All model mutations go through here."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_from_string


def _split_ref(ref: str) -> tuple[str, str]:
    """Split 'Sheet!B14' into ('Sheet', 'B14')."""
    if "!" not in ref:
        raise ValueError(f"Cell ref must be 'Sheet!Cell', got {ref!r}")
    sheet, cell = ref.split("!", 1)
    return sheet.strip("'"), cell


def read_cell(path: Path, ref: str) -> Any:
    """Read a single cell. Uses cached values (data_only=True)."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        sheet, cell = _split_ref(ref)
        return wb[sheet][cell].value
    finally:
        wb.close()


def read_cells(path: Path, refs: dict[str, str]) -> dict[str, Any]:
    """Read multiple cells in one open. refs = {name: 'Sheet!Cell'}."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        out: dict[str, Any] = {}
        for name, ref in refs.items():
            sheet, cell = _split_ref(ref)
            out[name] = wb[sheet][cell].value
        return out
    finally:
        wb.close()


@dataclass
class CellWrite:
    ref: str         # 'Sheet!B7'
    value: Any
    label: str = ""  # human-friendly label, e.g. "Q1-2026 toll revenue"


@dataclass
class CellWriteResult:
    ref: str
    label: str
    before: Any
    after: Any


def write_cells(path: Path, writes: list[CellWrite]) -> list[CellWriteResult]:
    """Apply a batch of cell writes. Returns before/after for the audit log.

    openpyxl writes raw cell values — it does NOT evaluate formulas. After a
    write, the CLI calls `fmcli.compute.recompute()` to update derived
    output cells in Python. Inputs change → outputs update; the .xlsx stays
    consistent without depending on Excel or LibreOffice for recalc.
    """
    wb = openpyxl.load_workbook(path)
    results: list[CellWriteResult] = []
    try:
        for w in writes:
            sheet, cell = _split_ref(w.ref)
            ws = wb[sheet]
            before = ws[cell].value
            ws[cell] = w.value
            results.append(
                CellWriteResult(ref=w.ref, label=w.label, before=before, after=w.value)
            )
        wb.save(path)
    finally:
        wb.close()
    return results


def list_named_outputs(path: Path, refs: dict[str, str]) -> list[tuple[str, str, Any]]:
    """Return [(name, ref, value)] for each output cell."""
    values = read_cells(path, refs)
    return [(name, refs[name], values[name]) for name in refs]


def read_borrower_submission(path: Path, sheet: str | None = None) -> dict[str, Any]:
    """Read a borrower-submitted workbook as a flat {label: value} dict.

    Convention: the workbook has a sheet (named in `sheet` or the first sheet)
    with two columns: A = label, B = value. Header row optional. This is the
    format we ask borrowers to submit; if they submit something else, the
    model-updater agent normalizes it before calling us.
    """
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb[sheet] if sheet else wb.worksheets[0]
        out: dict[str, Any] = {}
        for row in ws.iter_rows(values_only=True):
            if not row or len(row) < 2:
                continue
            label, value = row[0], row[1]
            if label is None or value is None:
                # Skip blank rows and descriptive note rows (label-only).
                continue
            label = str(label).strip()
            if not label or label.lower() in {"label", "metric", "field"}:
                continue
            out[label] = value
        return out
    finally:
        wb.close()


def workbook_summary(path: Path) -> dict[str, list[str]]:
    """Return {sheet_name: [first-col labels]} for orientation."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        out: dict[str, list[str]] = {}
        for ws in wb.worksheets:
            labels: list[str] = []
            for row in ws.iter_rows(min_row=1, max_row=40, max_col=1, values_only=True):
                v = row[0]
                if v is not None:
                    labels.append(str(v))
            out[ws.title] = labels
        return out
    finally:
        wb.close()


__all__ = [
    "CellWrite",
    "CellWriteResult",
    "read_cell",
    "read_cells",
    "write_cells",
    "list_named_outputs",
    "read_borrower_submission",
    "workbook_summary",
]
